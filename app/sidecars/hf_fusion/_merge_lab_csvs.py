# Merge fused tray ROI spectra with lab Brix/pH (offline). No hardware I/O.
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"D:\Data_JY\2026_Grape_Data_Collection\2026_Geneva_Concord")
XLSX = Path(r"D:\Data_JY\2026_Grape_Data_Collection\2026 Grape HSI Data Collection.xlsx")


def tray_sort_key(name: str):
    stage, _, rest = name.partition("_T")
    try:
        n = int(rest)
    except ValueError:
        n = 0
    return (0 if stage == "Unripe" else 1, n, name)


def load_lab() -> dict[tuple[str, int], tuple[object, object]]:
    wb = load_workbook(XLSX, data_only=True, read_only=True)
    ws = wb["Brix_pH_Data"]
    ref: dict[tuple[str, int], tuple[object, object]] = {}
    tray = None
    for row in ws.iter_rows(min_row=2, max_col=7, values_only=True):
        _date, _time, _cultivar, tray_name, sample, brix, ph = row[:7]
        if tray_name:
            tray = str(tray_name).strip()
        if tray is None or sample is None:
            continue
        try:
            sample_i = int(sample)
        except (TypeError, ValueError):
            continue
        ref[(tray, sample_i)] = (brix, ph)
    wb.close()
    return ref


def merge_mode(mode: str, ref: dict[tuple[str, int], tuple[object, object]]) -> None:
    rows = []
    header = None
    trays = []
    sessions = sorted(
        [p for p in ROOT.iterdir() if p.is_dir() and "_T" in p.name],
        key=lambda p: tray_sort_key(p.name),
    )
    for sess in sessions:
        csv_path = sess / mode / "fusion" / f"roi_spectra_fx10e_swir_{mode}.csv"
        if not csv_path.is_file():
            continue
        trays.append(sess.name)
        with csv_path.open(newline="", encoding="utf-8") as handle:
            reader = csv.DictReader(handle)
            if header is None:
                extra = [
                    c
                    for c in (reader.fieldnames or [])
                    if c not in {"image", "roi#", "label", "pixel_num", "Brix", "pH", "tray"}
                ]
                header = ["tray", "image", "roi#", "Brix", "pH", "pixel_num"] + extra
            for rec in reader:
                roi = int(rec["roi#"])
                brix, ph = ref.get((sess.name, roi), ("", ""))
                out = {
                    "tray": sess.name,
                    "image": rec.get("image", ""),
                    "roi#": roi,
                    "Brix": brix if brix is not None else "",
                    "pH": ph if ph is not None else "",
                    "pixel_num": rec.get("pixel_num", ""),
                }
                for col in header[6:]:
                    out[col] = rec.get(col, "")
                rows.append(out)
    dest = ROOT / f"roi_spectra_fx10e_swir_{mode}.csv"
    with dest.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=header)
        writer.writeheader()
        writer.writerows(rows)
    print(f"{mode}: {len(trays)} trays, {len(rows)} rows -> {dest}")
    print(" ", ", ".join(trays))


def main() -> int:
    ref = load_lab()
    for mode in ("reflectance", "transmittance"):
        merge_mode(mode, ref)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
